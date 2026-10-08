"""Regenerates the document fixtures from the vendored TeaStore text:
`uv run --with reportlab python services/extraction-agent/tests/fixtures/documents/make_fixtures.py`
(python-docx comes with docling). Then rewrite the golden files (SANDMAN_UPDATE_GOLDEN=1) and review the diff."""

from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph, SimpleDocTemplate

HERE = Path(__file__).parent
SENTENCES = (HERE.parent / "benchmark" / "teastore" / "teastore.txt").read_text(encoding="utf-8").splitlines()
PARAGRAPH = 6  # sentences per paragraph


def paragraphs() -> list[str]:
    return [" ".join(SENTENCES[i : i + PARAGRAPH]) for i in range(0, len(SENTENCES), PARAGRAPH)]


def teastore_pdf(path: Path) -> None:
    styles = getSampleStyleSheet()
    SimpleDocTemplate(str(path), pagesize=A4, title="TeaStore").build(
        [Paragraph(p.replace("&", "&amp;").replace("<", "&lt;"), styles["BodyText"]) for p in paragraphs()]
    )


def teastore_docx(path: Path) -> None:
    doc = Document()
    for p in paragraphs():
        doc.add_paragraph(p)
    doc.save(str(path))


def structure_docx(path: Path) -> None:
    doc = Document()
    doc.add_heading("Architecture", 1)
    doc.add_paragraph("The system has two components. The first stores media.")
    doc.add_heading("Components", 2)
    for item in ("The Cache keeps hot items.", "The Store persists all items."):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("Interfaces", 2)
    table = doc.add_table(rows=2, cols=2)
    for cell, text in zip(table.rows[0].cells, ("Name", "Provider"), strict=True):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ("IMediaStore", "Store"), strict=True):
        cell.text = text
    doc.add_paragraph("Clients use the interface. It is stable.")
    doc.save(str(path))


def selfcheck_pdf(path: Path) -> None:
    c = canvas.Canvas(str(path), pagesize=A4)
    c.drawString(72, 750, "The prefetch self-check converts this page offline.")
    c.save()


if __name__ == "__main__":
    teastore_pdf(HERE / "teastore.pdf")
    teastore_docx(HERE / "teastore.docx")
    structure_docx(HERE / "structure.docx")
    selfcheck_pdf(HERE.parents[2] / "src" / "sandman_extraction" / "selfcheck.pdf")
